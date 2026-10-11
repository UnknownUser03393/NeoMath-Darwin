from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
PPT = ROOT / "submission" / "NeoMath_作品说明.pptx"


def main() -> int:
    if not PPT.exists():
        raise FileNotFoundError(PPT)
    presentation = Presentation(PPT)
    issues: list[str] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0:
                issues.append(f"slide {slide_index}: {shape.name} starts outside the canvas")
            if shape.left + shape.width > presentation.slide_width:
                issues.append(f"slide {slide_index}: {shape.name} exceeds right edge")
            if shape.top + shape.height > presentation.slide_height:
                issues.append(f"slide {slide_index}: {shape.name} exceeds bottom edge")
    report = {
        "ok": len(presentation.slides) == 16 and not issues,
        "path": str(PPT),
        "slideCount": len(presentation.slides),
        "issues": issues,
        "sizeBytes": PPT.stat().st_size,
    }
    output = ROOT / "build" / "reports" / "ppt_validation.json"
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
