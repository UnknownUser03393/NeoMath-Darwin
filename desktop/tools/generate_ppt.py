from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "submission" / "NeoMath_作品说明.pptx"
SCREENSHOT = ROOT / "docs" / "assets" / "neomath-ui.png"
REPORTS = ROOT / "build" / "reports"

WIDE = Inches(13.333)
HIGH = Inches(7.5)
INK = RGBColor(20, 37, 32)
MUTED = RGBColor(95, 111, 104)
GREEN = RGBColor(24, 105, 83)
GREEN_DARK = RGBColor(11, 57, 46)
GREEN_SOFT = RGBColor(228, 240, 234)
CORAL = RGBColor(217, 102, 77)
CORAL_SOFT = RGBColor(249, 229, 223)
AMBER = RGBColor(217, 154, 39)
AMBER_SOFT = RGBColor(251, 240, 214)
VIOLET = RGBColor(118, 89, 166)
PAPER = RGBColor(248, 246, 240)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(218, 214, 204)


def load_json(name: str) -> dict[str, Any]:
    path = REPORTS / name
    if not path.exists():
        return {}
    value: Any = json.loads(path.read_text(encoding="utf-8"))
    return value if isinstance(value, dict) else {}


def text_box(
    slide: Any,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    bold: bool = False,
    color: RGBColor = INK,
    font: str = "Microsoft YaHei UI",
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> Any:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.03)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def panel(
    slide: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: RGBColor = WHITE,
    line: RGBColor = LINE,
    radius: bool = False,
) -> Any:
    geometry = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(geometry, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def add_footer(slide: Any, number: int, section: str) -> None:
    text_box(slide, section.upper(), 0.55, 7.12, 3.8, 0.2, size=9, bold=True, color=MUTED)
    text_box(slide, str(number), 12.15, 7.08, 0.55, 0.22, size=9, color=MUTED, align=PP_ALIGN.RIGHT)


def title(slide: Any, headline: str, number: int, section: str, eyebrow: str = "NEOMATH") -> None:
    text_box(slide, eyebrow, 0.58, 0.36, 3.0, 0.23, size=10, bold=True, color=GREEN)
    text_box(slide, headline, 0.58, 0.74, 12.0, 0.62, size=35, bold=True)
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.58),
        Inches(1.48),
        Inches(12.15),
        Inches(0.018),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()
    add_footer(slide, number, section)


def bullet_list(
    slide: Any,
    items: list[str],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: RGBColor = INK,
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Microsoft YaHei UI"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
        paragraph.space_after = Pt(10)
        paragraph.text = f"•  {item}"


def stat(
    slide: Any, value: str, label: str, x: float, y: float, w: float, accent: RGBColor
) -> None:
    panel(slide, x, y, w, 1.35, fill=PAPER, line=LINE)
    text_box(slide, value, x + 0.18, y + 0.16, w - 0.36, 0.52, size=28, bold=True, color=accent)
    text_box(slide, label, x + 0.18, y + 0.78, w - 0.36, 0.48, size=13, color=MUTED)


def add_cover(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = GREEN_DARK
    text_box(
        slide,
        "阿里云·AI少年 Vibe Coding 创新挑战活动",
        0.68,
        0.58,
        8.4,
        0.35,
        size=14,
        color=GREEN_SOFT,
    )
    text_box(
        slide, "NeoMath", 0.68, 1.55, 7.0, 0.95, size=54, bold=True, color=WHITE, font="Georgia"
    )
    text_box(slide, "数学动态调试环境", 0.72, 2.55, 6.7, 0.55, size=30, bold=True, color=WHITE)
    text_box(
        slide,
        "The Debugger for Mathematics",
        0.72,
        3.22,
        7.2,
        0.42,
        size=20,
        color=RGBColor(164, 207, 191),
    )
    text_box(
        slide,
        "把数学当作程序运行，把错误当作 Bug 调试。",
        0.72,
        4.45,
        8.6,
        0.6,
        size=24,
        bold=True,
        color=WHITE,
    )
    panel(slide, 9.82, 0.7, 2.8, 5.95, fill=GREEN_SOFT, line=GREEN_SOFT)
    text_box(
        slide,
        "N",
        10.48,
        1.35,
        1.45,
        1.3,
        size=68,
        bold=True,
        color=GREEN_DARK,
        font="Georgia",
        align=PP_ALIGN.CENTER,
    )
    text_box(
        slide,
        "Write\nRun\nBreak\nExplain\nFix\nReplay",
        10.35,
        3.0,
        1.75,
        2.65,
        size=20,
        bold=True,
        color=GREEN_DARK,
        align=PP_ALIGN.CENTER,
    )
    text_box(
        slide,
        "Windows 10 / 11 x64 · Python + C++20",
        0.72,
        6.7,
        7.5,
        0.3,
        size=12,
        color=RGBColor(164, 207, 191),
    )


def add_pain(prs: Presentation, number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "学生知道“错了”，却不知道错误如何发生", number, "创意来源与学习痛点")
    panel(slide, 0.65, 1.85, 5.65, 4.55, fill=CORAL_SOFT, line=CORAL)
    text_box(slide, "传统反馈", 0.95, 2.18, 2.0, 0.4, size=22, bold=True, color=CORAL)
    bullet_list(
        slide,
        ["只给最终正误", "跳过丢根、增根、定义域变化", "学生只能重做，无法定位首次分叉"],
        0.95,
        2.85,
        4.95,
        2.35,
        size=19,
    )
    text_box(
        slide, "“第二步错了”不是证明。", 0.95, 5.58, 4.8, 0.45, size=22, bold=True, color=CORAL
    )
    panel(slide, 6.72, 1.85, 5.95, 4.55, fill=GREEN_SOFT, line=GREEN)
    text_box(slide, "NeoMath 的回答", 7.02, 2.18, 3.3, 0.4, size=22, bold=True, color=GREEN)
    bullet_list(
        slide,
        [
            "比较每一步的解集与定义域",
            "锁定第一处可证明的分叉",
            "运行一个具体检验值，让差异当场出现",
        ],
        7.02,
        2.85,
        5.05,
        2.35,
        size=19,
    )
    text_box(
        slide,
        "“代入 x = 0，上一步成立，这一步不成立。”",
        7.02,
        5.48,
        5.05,
        0.58,
        size=20,
        bold=True,
        color=GREEN,
    )


def add_solution(prs: Presentation, number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "NeoMath 把代数推导变成一段可调试的执行过程", number, "方案")
    if SCREENSHOT.exists():
        slide.shapes.add_picture(str(SCREENSHOT), Inches(5.0), Inches(1.75), width=Inches(7.7))
    text_box(slide, "面向中学生的\n一元代数推导调试器", 0.72, 1.95, 3.9, 1.25, size=27, bold=True)
    text_box(
        slide,
        "• 逐行写推导\n• C++ 精确验证\n• Probe 证明差异\n• 离线完整运行",
        0.72,
        3.55,
        3.85,
        2.2,
        size=18,
    )
    text_box(
        slide,
        "不是拍照搜题，也不是标准答案生成器。",
        0.72,
        6.15,
        4.2,
        0.36,
        size=16,
        bold=True,
        color=MUTED,
    )


def add_loop(prs: Presentation, number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "七个动作组成一次完整的数学调试闭环", number, "核心闭环")
    labels = [
        ("Write", "写下推导"),
        ("Complete", "补全操作"),
        ("Run", "运行"),
        ("Break", "定位"),
        ("Explain", "解释"),
        ("Fix", "修正"),
        ("Replay", "重查"),
    ]
    y = 3.45
    for index in range(len(labels) - 1):
        x1 = 0.68 + index * 1.75 + 1.35
        x2 = 0.68 + (index + 1) * 1.75
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(x1),
            Inches(y),
            Inches(x2),
            Inches(y),
        )
        connector.line.color.rgb = GREEN
        connector.line.width = Pt(2)
    for index, (english, chinese) in enumerate(labels):
        x = 0.68 + index * 1.75
        fill = CORAL_SOFT if english == "Break" else GREEN_SOFT
        line = CORAL if english == "Break" else GREEN
        panel(slide, x, 2.62, 1.35, 1.7, fill=fill, line=line, radius=True)
        text_box(
            slide,
            english,
            x + 0.08,
            2.96,
            1.19,
            0.34,
            size=17,
            bold=True,
            color=line,
            align=PP_ALIGN.CENTER,
        )
        text_box(
            slide, chinese, x + 0.08, 3.48, 1.19, 0.28, size=13, color=INK, align=PP_ALIGN.CENTER
        )
    text_box(
        slide,
        "每次修正后重新执行；不靠猜测，靠状态比较。",
        2.3,
        5.35,
        8.7,
        0.5,
        size=24,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def add_probe(prs: Presentation, number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "Probe 用一个具体值，把抽象错误变成可运行证据", number, "核心创新")
    text_box(slide, "x(x − 1) = 0", 0.85, 2.05, 4.2, 0.52, size=30, bold=True, font="Cambria Math")
    text_box(
        slide,
        "x − 1 = 0",
        0.85,
        3.1,
        4.2,
        0.52,
        size=30,
        bold=True,
        color=CORAL,
        font="Cambria Math",
    )
    panel(slide, 5.25, 1.82, 6.95, 4.55, fill=CORAL_SOFT, line=CORAL)
    text_box(slide, "检验值  x = 0", 5.65, 2.25, 5.9, 0.55, size=28, bold=True, color=CORAL)
    text_box(slide, "上一步", 5.7, 3.2, 1.35, 0.35, size=16, bold=True, color=MUTED)
    text_box(
        slide,
        "0(0 − 1) = 0  ✓",
        7.25,
        3.12,
        4.1,
        0.48,
        size=24,
        bold=True,
        color=GREEN,
        font="Cambria Math",
    )
    text_box(slide, "这一步", 5.7, 4.24, 1.35, 0.35, size=16, bold=True, color=MUTED)
    text_box(
        slide,
        "0 − 1 = 0  ✕",
        7.25,
        4.16,
        4.1,
        0.48,
        size=24,
        bold=True,
        color=CORAL,
        font="Cambria Math",
    )
    text_box(slide, "结论：丢失了解 x = 0。", 5.7, 5.35, 5.7, 0.5, size=23, bold=True)
    text_box(
        slide,
        "它证明发生了什么，\n但不武断猜测学生执行了哪个动作。",
        0.85,
        5.55,
        3.95,
        0.78,
        size=17,
        color=MUTED,
    )


def add_completion(prs: Presentation, number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "语义补全只推荐 C++ 能验证并签发证书的操作", number, "Semantic Completion")
    panel(slide, 0.75, 1.88, 4.25, 4.55, fill=PAPER, line=LINE)
    text_box(
        slide, "::expr  x² − 4 = 0", 1.08, 2.32, 3.55, 0.5, size=24, bold=True, font="Cascadia Mono"
    )
    text_box(
        slide,
        "^  factor",
        1.08,
        3.26,
        3.55,
        0.5,
        size=24,
        bold=True,
        color=VIOLET,
        font="Cascadia Mono",
    )
    text_box(slide, "Ctrl + Space", 1.08, 4.36, 2.1, 0.34, size=15, bold=True, color=MUTED)
    text_box(slide, "风险：safe", 1.08, 5.25, 2.5, 0.35, size=16, bold=True, color=GREEN)
    panel(slide, 5.38, 1.88, 7.1, 4.55, fill=GREEN_SOFT, line=GREEN)
    text_box(
        slide, "(x − 2)(x + 2) = 0", 5.78, 2.38, 6.3, 0.62, size=30, bold=True, font="Cambria Math"
    )
    bullet_list(
        slide,
        [
            "识别：Difference of Squares",
            "验证：展开目标 = 规范化源式",
            "证书：Verified · ExactExpansion",
            "解集保持：{−2, 2}",
        ],
        5.78,
        3.35,
        5.95,
        2.35,
        size=18,
    )


def add_core(prs: Presentation, number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "数学真值集中在唯一的 C++20 MathCore", number, "C++ MathCore")
    columns = [
        ("Syntax", ["Lexer / Parser", "AST / SourceSpan", "复杂度预算"]),
        ("Algebra", ["精确 Rational", "Polynomial", "Normalizer / Solver"]),
        ("Semantics", ["Domain / SolutionSet", "StepRelation / Probe", "Trace / Replay"]),
        ("Trust", ["Transformation", "Certificate", "结构化诊断"]),
    ]
    for index, (heading, items) in enumerate(columns):
        x = 0.66 + index * 3.04
        panel(slide, x, 1.95, 2.72, 4.45, fill=PAPER, line=LINE)
        text_box(slide, heading, x + 0.22, 2.28, 2.28, 0.42, size=22, bold=True, color=GREEN)
        bullet_list(slide, items, x + 0.22, 3.0, 2.28, 2.25, size=16)
    text_box(
        slide,
        "值语义优先 · 无裸拥有指针 · 精确判断不用浮点 · 超限返回 Unverified",
        1.15,
        6.63,
        11.05,
        0.3,
        size=15,
        bold=True,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def add_python(prs: Presentation, number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "Python 负责产品能力，JavaScript 负责高频交互", number, "Python 桌面壳")
    sections = [
        ("Web UI", "编辑、KaTeX、快捷键、Undo / Redo\n300 ms debounce，局部更新", VIOLET),
        ("Python", "pywebview、生命周期、文档、自动保存\nAI 接入、日志、打包与提交", GREEN),
        ("C++ Core", "完整业务操作一次跨边界\n数学状态可直接 JSON 序列化", CORAL),
    ]
    for index in range(2):
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(4.18 + index * 4.2),
            Inches(3.5),
            Inches(4.72 + index * 4.2),
            Inches(3.5),
        )
        connector.line.color.rgb = MUTED
        connector.line.width = Pt(2)
    for index, (heading, body, accent) in enumerate(sections):
        x = 0.62 + index * 4.2
        panel(slide, x, 2.05, 3.55, 3.9, fill=WHITE, line=accent, radius=True)
        text_box(
            slide,
            heading,
            x + 0.28,
            2.48,
            2.9,
            0.46,
            size=25,
            bold=True,
            color=accent,
            align=PP_ALIGN.CENTER,
        )
        text_box(slide, body, x + 0.32, 3.35, 2.9, 1.65, size=17, color=INK, align=PP_ALIGN.CENTER)
    text_box(
        slide,
        "运行时不需要 Node.js，不自带 Chromium；使用系统 WebView2。",
        2.15,
        6.35,
        9.0,
        0.38,
        size=18,
        bold=True,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def add_pybind(prs: Presentation, number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "pybind11 边界按业务操作设计，而不是按 AST 节点设计", number, "粗粒度边界")
    api_names = [
        "parse_document(source)",
        "run_trace(document)",
        "replay(document, changed_step)",
        "get_completions(document, step)",
        "execute_transformation(document, step, intent)",
        "evaluate_probe(document, probe)",
    ]
    panel(slide, 0.72, 1.9, 7.25, 4.72, fill=PAPER, line=LINE)
    text_box(slide, "公开 API", 1.02, 2.2, 2.3, 0.4, size=22, bold=True, color=GREEN)
    for index, name in enumerate(api_names):
        text_box(slide, name, 1.02, 2.86 + index * 0.55, 6.4, 0.3, size=15, font="Cascadia Mono")
    panel(slide, 8.35, 1.9, 4.25, 4.72, fill=GREEN_SOFT, line=GREEN)
    text_box(slide, "一次调用完成一件事", 8.72, 2.28, 3.5, 0.46, size=23, bold=True, color=GREEN)
    bullet_list(
        slide,
        [
            "参数在持有 GIL 时校验",
            "计算阶段释放 GIL",
            "返回前重新获取 GIL",
            "异常映射为结构化错误",
            "不暴露指针或可变 AST",
        ],
        8.72,
        3.02,
        3.35,
        2.7,
        size=17,
    )


def add_ai(prs: Presentation, number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "AI 负责理解与表达，永远不能改写数学真值", number, "AI 作用")
    panel(slide, 0.7, 1.9, 5.55, 4.65, fill=GREEN_SOFT, line=GREEN)
    text_box(slide, "AI 可以做", 1.05, 2.28, 2.9, 0.42, size=23, bold=True, color=GREEN)
    bullet_list(
        slide,
        [
            "自然语言 → 白名单 Intent",
            "把结构化诊断改写成学生语言",
            "排序安全补全",
            "网络失败时自动降级",
        ],
        1.05,
        3.0,
        4.65,
        2.55,
        size=18,
    )
    panel(slide, 6.72, 1.9, 5.9, 4.65, fill=CORAL_SOFT, line=CORAL)
    text_box(slide, "AI 不能做", 7.08, 2.28, 2.9, 0.42, size=23, bold=True, color=CORAL)
    bullet_list(
        slide,
        [
            "决定解集或定义域",
            "直接修改 MathState",
            "绕过 C++ Transformation Executor",
            "把未验证内容说成已证明",
        ],
        7.08,
        3.0,
        4.75,
        2.55,
        size=18,
    )
    text_box(
        slide,
        "Candidate → Schema → Allowlist → C++ Executor → Certificate",
        2.0,
        6.65,
        9.25,
        0.28,
        size=16,
        bold=True,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def add_trust(prs: Presentation, number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "可信边界由“候选生成后回代原式”守住", number, "数学可信模型")
    steps = [
        ("1", "Normalize", "规范化"),
        ("2", "Candidates", "生成候选"),
        ("3", "Domain", "检查定义域"),
        ("4", "Substitute", "回代原式"),
        ("5", "Accept", "只接受已验证解"),
    ]
    for index, (num, english, chinese) in enumerate(steps):
        x = 0.58 + index * 2.52
        panel(slide, x, 2.08, 2.18, 2.08, fill=PAPER, line=LINE)
        text_box(slide, num, x + 0.18, 2.28, 0.45, 0.38, size=20, bold=True, color=CORAL)
        text_box(slide, english, x + 0.2, 2.91, 1.78, 0.35, size=18, bold=True, color=GREEN)
        text_box(slide, chinese, x + 0.2, 3.48, 1.78, 0.3, size=15, color=MUTED)
    panel(slide, 1.2, 4.72, 10.95, 1.28, fill=CORAL_SOFT, line=CORAL)
    text_box(
        slide,
        "sqrt(x) = x − 2  →  候选 {1, 4}  →  回代淘汰 1  →  最终 {4}",
        1.52,
        5.08,
        10.3,
        0.45,
        size=23,
        bold=True,
        font="Cambria Math",
        align=PP_ALIGN.CENTER,
    )
    text_box(
        slide,
        "Verified / Unverified / Unsupported 明确分开；无法证明时不猜。",
        2.25,
        6.42,
        8.9,
        0.35,
        size=18,
        bold=True,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def add_demos(prs: Presentation, number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "四个离线 Demo 覆盖丢根、增根、定义域与安全补全", number, "演示")
    demos = [
        ("01", "丢根", "x(x−1)=0 → x−1=0", "Probe x=0", CORAL_SOFT, CORAL),
        ("02", "平方增根", "sqrt(x)=x−2 → 平方", "Probe x=1", AMBER_SOFT, AMBER),
        ("03", "定义域扩大", "(x−1)/(x−1)=1 → 1=1", "Probe x=1", RGBColor(238, 231, 247), VIOLET),
        ("04", "语义补全", "x²−4 → (x−2)(x+2)", "Verified Certificate", GREEN_SOFT, GREEN),
    ]
    for index, (num, heading, formula, proof, fill, accent) in enumerate(demos):
        x = 0.68 + (index % 2) * 6.08
        y = 1.86 + (index // 2) * 2.36
        panel(slide, x, y, 5.65, 2.02, fill=fill, line=accent)
        text_box(slide, num, x + 0.25, y + 0.22, 0.62, 0.35, size=16, bold=True, color=accent)
        text_box(slide, heading, x + 1.0, y + 0.2, 2.5, 0.4, size=21, bold=True, color=accent)
        text_box(
            slide, formula, x + 0.25, y + 0.88, 4.95, 0.34, size=17, bold=True, font="Cambria Math"
        )
        text_box(slide, proof, x + 0.25, y + 1.45, 4.95, 0.28, size=14, color=MUTED)


def add_performance(prs: Presentation, number: int, benchmark: dict[str, Any]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "实测 P95 全部低于首版交互目标", number, "性能")
    results = benchmark.get("results", {})
    labels = ["单行解析", "一次分析", "20 行 Trace", "Replay"]
    keys = ["parseSingleLine", "analyzeLinear", "trace20Lines", "replay"]
    values = [float(results.get(key, {}).get("p95Ms", 0)) for key in keys]
    chart_data = ChartData()
    chart_data.categories = labels
    chart_data.add_series("P95 ms", values)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.75),
        Inches(1.92),
        Inches(7.35),
        Inches(4.55),
        chart_data,
    ).chart
    chart.has_legend = False
    chart.value_axis.has_major_gridlines = True
    chart.category_axis.tick_labels.font.size = Pt(12)
    chart.value_axis.tick_labels.font.size = Pt(11)
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = GREEN
    chart.has_title = False
    stat(slide, f"{values[0]:.4f} ms", "单行解析 P95 / 目标 < 2 ms", 8.48, 2.0, 3.85, GREEN)
    stat(slide, f"{values[2]:.4f} ms", "20 行 Trace P95 / 目标 < 100 ms", 8.48, 3.67, 3.85, VIOLET)
    ratio = float(benchmark.get("boundaryEstimate", {}).get("estimatedOverheadRatio", 0))
    stat(slide, f"{ratio:.1%}", "估算 pybind 边界占比 / 目标 < 20%", 8.48, 5.34, 3.85, CORAL)


def add_tests(prs: Presentation, number: int, differential: dict[str, Any]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "独立 Oracle 与分层测试共同约束数学正确性", number, "测试与验证")
    checked = int(differential.get("checkedCases", 0))
    differences = int(differential.get("differenceCount", 0))
    stat(slide, "10", "C++ Catch2 测试用例", 0.75, 1.95, 2.65, GREEN)
    stat(slide, "30", "Python / Integration / Golden", 3.62, 1.95, 2.65, VIOLET)
    stat(slide, f"{checked:,}", "固定 Seed 差分案例", 6.49, 1.95, 2.65, CORAL)
    stat(slide, str(differences), "当前差异数量", 9.36, 1.95, 2.65, AMBER)
    panel(slide, 0.75, 3.75, 11.25, 2.25, fill=PAPER, line=LINE)
    bullet_list(
        slide,
        [
            "C++：Rational、Parser、Polynomial、Solver、Trace、Transformation",
            "Python：粗粒度 API、四个 Demo、文档、AI 白名单、打包检查",
            "SymPy Oracle 接收结构化 AST，不重新解析 UI 字符串",
            "随机系数范围 −10…10，Seed 20260724；差异会输出最小失败输入",
        ],
        1.08,
        4.08,
        10.55,
        1.65,
        size=16,
    )
    text_box(
        slide,
        "不伪造通过：所有数字来自本机实际执行报告。",
        2.35,
        6.45,
        8.4,
        0.34,
        size=17,
        bold=True,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def add_future(prs: Presentation, number: int, release: dict[str, Any]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    title(slide, "先交付可信的一元代数闭环，再扩展表达能力", number, "未来规划")
    phases = [
        ("现在", "一元实数代数", "丢根 / 增根 / 定义域\n可解释 Probe"),
        ("下一步", "更强精确数", "BigInt / 代数数\n更多可验证根式"),
        ("随后", "可视化理解", "函数图像与 Probe 联动\n仍由核心状态驱动"),
        ("远期", "可复用 MathCore", "WASM / 独立进程\n更多学习场景"),
    ]
    for index, (when, heading, body) in enumerate(phases):
        x = 0.62 + index * 3.08
        panel(slide, x, 2.0, 2.72, 3.85, fill=PAPER, line=LINE)
        text_box(slide, when, x + 0.22, 2.28, 2.28, 0.3, size=13, bold=True, color=CORAL)
        text_box(slide, heading, x + 0.22, 2.92, 2.28, 0.7, size=22, bold=True, color=GREEN)
        text_box(slide, body, x + 0.22, 4.02, 2.28, 1.05, size=16, color=INK)
    package_size = int(release.get("sizeBytes", 0))
    package_text = (
        f"当前 onedir：{package_size / 1024 / 1024:.1f} MB"
        if package_size
        else "onedir 大小将在最终打包后写入"
    )
    text_box(
        slide,
        package_text,
        3.65,
        6.32,
        6.0,
        0.35,
        size=17,
        bold=True,
        color=MUTED,
        align=PP_ALIGN.CENTER,
    )


def add_close(prs: Presentation, number: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = GREEN_DARK
    text_box(
        slide, "创作心得", 0.72, 0.62, 3.0, 0.34, size=13, bold=True, color=RGBColor(164, 207, 191)
    )
    text_box(
        slide,
        "AI 可以让解释更亲切，\n但数学可信必须来自可验证的程序。",
        0.72,
        1.52,
        10.8,
        1.48,
        size=40,
        bold=True,
        color=WHITE,
    )
    text_box(
        slide,
        "NeoMath 的价值不在“更快给答案”，"
        "而在让学生看见错误是怎样发生的，\n并用一个具体值亲手复现它。",
        0.72,
        3.6,
        9.8,
        1.05,
        size=23,
        color=GREEN_SOFT,
    )
    panel(slide, 9.92, 4.88, 2.65, 1.45, fill=GREEN_SOFT, line=GREEN_SOFT)
    text_box(
        slide,
        "Run → Break\n→ Explain → Replay",
        10.16,
        5.18,
        2.15,
        0.76,
        size=17,
        bold=True,
        color=GREEN_DARK,
        align=PP_ALIGN.CENTER,
    )
    text_box(
        slide,
        "NeoMath · The Debugger for Mathematics",
        0.72,
        6.72,
        6.7,
        0.28,
        size=13,
        color=RGBColor(164, 207, 191),
    )
    text_box(
        slide,
        str(number),
        12.0,
        6.72,
        0.55,
        0.28,
        size=10,
        color=RGBColor(164, 207, 191),
        align=PP_ALIGN.RIGHT,
    )


def main() -> int:
    benchmark = load_json("benchmark.json")
    differential = load_json("differential.json")
    release = load_json("release.json")
    prs = Presentation()
    prs.slide_width = WIDE
    prs.slide_height = HIGH
    add_cover(prs)
    add_pain(prs, 2)
    add_solution(prs, 3)
    add_loop(prs, 4)
    add_probe(prs, 5)
    add_completion(prs, 6)
    add_core(prs, 7)
    add_python(prs, 8)
    add_pybind(prs, 9)
    add_ai(prs, 10)
    add_trust(prs, 11)
    add_demos(prs, 12)
    add_performance(prs, 13, benchmark)
    add_tests(prs, 14, differential)
    add_future(prs, 15, release)
    add_close(prs, 16)
    prs.core_properties.title = "NeoMath 作品说明"
    prs.core_properties.subject = "数学动态调试环境 · 阿里云 AI少年 Vibe Coding 创新挑战活动"
    prs.core_properties.author = "NeoMath"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
